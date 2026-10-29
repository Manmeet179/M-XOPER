import io
import os
import re
import zipfile
import mimetypes
from datetime import date, datetime

import streamlit as st
import openpyxl
import fitz
from PIL import Image
try:
    import psycopg2
except ImportError:
    psycopg2 = None
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


# ============================================================
# APP CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="M-CERTIFICATE",
    page_icon="〽️",
    layout="wide",
)

DEFAULTS = {
    "date_format": "DD-MMM-YY",
    "company": "OptoPixel Pvt. Ltd.",
    "paper_type": "Letterhead",
    "top_spacing_mm": 10,
    "show_make": True,
    "show_model": True,
    "show_serial": True,
    "show_mac": True,
    "show_win": True,
    "show_operating_system": True,
    "make_heading": "MAKE",
    "model_heading": "MODEL NO",
    "serial_heading": "SERIAL NO",
    "mac_heading": "MAC ID",
    "win_heading": "WINDOW KEY",
    "operating_system_heading": "OPERATING SYSTEM",
    "table_order": ["item", "model_no", "serial_no", "operating_system", "win_key", "mac", "make"],
    "intro_model": True,
    "intro_serial": True,
    "intro_mac": False,
    "intro_make": False,
    "intro_operating_system": False,
    "intro_win": False,
}

DATE_FORMATS = {
    "DD-MMM-YY": "%d-%b-%y",
    "DD-MMM-YYYY": "%d-%b-%Y",
    "DD/MM/YYYY": "%d/%m/%Y",
    "DD-MM-YYYY": "%d-%m-%Y",
    "MM/DD/YYYY": "%m/%d/%Y",
    "YYYY-MM-DD": "%Y-%m-%d",
}

EMPTY_VALUES = {
    "", "n/a", "na", "n.a.", "not applicable",
    "-", "--", "—", "none", "null",
}

if "smart_settings" not in st.session_state:
    st.session_state.smart_settings = DEFAULTS.copy()

for key, value in DEFAULTS.items():
    st.session_state.smart_settings.setdefault(key, value)

st.session_state.setdefault("page", "M-SMART PDF")
st.session_state.setdefault("smart_generated", None)
st.session_state.setdefault("manual_generated", None)
st.session_state.setdefault("authenticated", False)
st.session_state.setdefault("is_admin", False)
st.session_state.setdefault("excel_upload_signature", None)
st.session_state.setdefault("windows_queue", [])
st.session_state.setdefault("windows_decisions", {})


# ============================================================
# PREMIUM UI
# ============================================================

st.markdown("""
<style>
.stApp {
    background: #000000 !important;
    color: #FFFFFF !important;
}
.block-container {
    max-width: 1500px;
    padding-top: 3.2rem;
}
.brand-title {
    text-align: center;
    font-size: 34px;
    font-weight: 850;
    letter-spacing: 2px;
    color: #FFFFFF !important;
}
.brand-subtitle {
    text-align: center;
    color: #FFFFFF !important;
    margin-bottom: 22px;
}
div[data-testid="stButton"] button {
    min-height: 45px;
    border-radius: 12px;
    font-weight: 700;
}
div[data-testid="stDownloadButton"] button {
    border-radius: 10px;
    font-weight: 700;
}
div[data-testid="stFileUploader"] {
    border-radius: 12px;
}
.stApp, .stApp p, .stApp span, .stApp label,
.stApp h1, .stApp h2, .stApp h3, .stApp h4,
.stApp h5, .stApp h6, [data-testid="stMarkdownContainer"],
[data-testid="stWidgetLabel"], [data-testid="stCaptionContainer"] {
    color: #FFFFFF !important;
}
.stApp input, .stApp textarea,
.stApp [data-baseweb="select"] {
    color: #FFFFFF !important;
    -webkit-text-fill-color: #FFFFFF !important;
}
</style>
""", unsafe_allow_html=True)

st.markdown(
    '<div class="brand-title">M-CERTIFICATE</div>'
    '<div class="brand-subtitle">'
    'SMART CERTIFICATE STUDIO · EXCEL TO PDF + WORD'
    '</div>',
    unsafe_allow_html=True,
)


# ============================================================
# LOGIN
# ============================================================

if not st.session_state.authenticated:
    left, center, right = st.columns([1, 1.2, 1])

    with center:
        st.subheader("Login to Continue")

        with st.form("secure_login_form"):
            username = st.text_input(
                "User ID",
                placeholder="Enter your ID",
            )
            password = st.text_input(
                "Password",
                type="password",
                placeholder="Enter your password",
            )
            login_clicked = st.form_submit_button(
                "LOGIN",
                type="primary",
                use_container_width=True,
            )

        if login_clicked:
            try:
                normal_user = (
                    username == str(st.secrets["login"]["username"])
                    and password == str(st.secrets["login"]["password"])
                )

                admin_user = (
                    username == str(st.secrets["admin"]["username"])
                    and password == str(st.secrets["admin"]["password"])
                )

                if admin_user:
                    st.session_state.authenticated = True
                    st.session_state.is_admin = True
                    st.rerun()

                elif normal_user:
                    st.session_state.authenticated = True
                    st.session_state.is_admin = False
                    st.rerun()

                else:
                    st.error("Invalid ID or password.")

            except Exception:
                st.error(
                    "Secrets configuration is missing. "
                    "Check your Streamlit Secrets settings."
                )

    st.markdown("""
    <div style="position:fixed;bottom:15px;left:0;width:100%;
    text-align:center;color:#FFFFFF;font-size:12px;letter-spacing:.5px">
    M-CERTIFICATE · SMART CERTIFICATE STUDIO<br>
    MADE BY ~ MEET MEVADA
    </div>
    """, unsafe_allow_html=True)

    st.stop()


# ============================================================
# AIVEN POSTGRESQL — PRIVATE LETTERHEAD IMAGE
# ============================================================

def get_db_connection():
    if psycopg2 is None:
        raise RuntimeError("psycopg2 is not installed. Install psycopg2-binary for database/letterhead features.")
    return psycopg2.connect(st.secrets["postgres"]["url"])


def init_logo_table():
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS private_assets (
                    name TEXT PRIMARY KEY,
                    mime_type TEXT NOT NULL,
                    content BYTEA NOT NULL,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
            """)


@st.cache_data(ttl=60)
def get_bg_logo():
    init_logo_table()

    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT content, mime_type
                FROM private_assets
                WHERE name = %s
                """,
                ("certificate_background",),
            )
            row = cur.fetchone()

    if not row:
        return None

    return bytes(row[0]), row[1]


def save_bg_logo(uploaded_file):
    if not st.session_state.get("is_admin", False):
        raise PermissionError("Administrator login is required.")

    filename = uploaded_file.name.lower()

    if not filename.endswith((".png", ".jpg", ".jpeg")):
        raise ValueError("Only PNG and JPG images are supported.")

    image_bytes = uploaded_file.getvalue()

    if not image_bytes:
        raise ValueError("The uploaded image is empty.")

    try:
        with Image.open(io.BytesIO(image_bytes)) as im:
            im.verify()
    except Exception as exc:
        raise ValueError(
            "The uploaded file is not a valid image."
        ) from exc

    mime_type = (
        mimetypes.guess_type(filename)[0]
        or "application/octet-stream"
    )

    init_logo_table()

    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO private_assets
                    (name, mime_type, content, updated_at)
                VALUES (%s, %s, %s, CURRENT_TIMESTAMP)
                ON CONFLICT (name) DO UPDATE SET
                    mime_type = EXCLUDED.mime_type,
                    content = EXCLUDED.content,
                    updated_at = CURRENT_TIMESTAMP
            """, (
                "certificate_background",
                mime_type,
                psycopg2.Binary(image_bytes),
            ))

    get_bg_logo.clear()


# ============================================================
# COMMON HELPERS
# ============================================================

def clean(value):
    return "" if value is None else str(value).strip()


def valid_value(value):
    value = clean(value)
    return "" if value.casefold() in EMPTY_VALUES else value


def safe_stem(value):
    value = re.sub(r"[^A-Za-z0-9.-]+", "_", str(value or ""))
    return value.strip("._-")[:100] or "TEMP"


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = clean(value)

    for fmt in (
        "%d-%b-%y",
        "%d-%b-%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    return date.today()


def item_filename(item):
    parts = [
        "TEST_CERTIFICATE",
        safe_stem(valid_value(item.get("item")) or "TEMP").upper(),
    ]

    model = valid_value(item.get("model_no"))
    serial = valid_value(item.get("serial_no"))

    if model:
        parts.append(safe_stem(model).upper())

    if serial:
        parts.append(safe_stem(serial).upper())

    return "_".join(parts)


def combined_filename(company, cert_date):
    company_part = safe_stem(company).upper()

    date_part = (
        cert_date.strftime("%d-%b-%y").upper()
        if isinstance(cert_date, (date, datetime))
        else safe_stem(cert_date).upper()
    )

    return f"TEST_CERTIFICATE_{company_part}_{date_part}"


def merged_cell_value(ws, row, column):
    if not column:
        return None

    cell = ws.cell(row, column)

    if cell.value is not None:
        return cell.value

    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return ws.cell(
                merged_range.min_row,
                merged_range.min_col,
            ).value

    return None


def wrap_text(text, font, size, width):
    words = str(text or "").split()

    if not words:
        return [""]

    lines = []
    current = ""

    for word in words:
        candidate = word if not current else current + " " + word

        if stringWidth(candidate, font, size) <= width:
            current = candidate
        else:
            if current:
                lines.append(current)

            current = word

            if stringWidth(current, font, size) > width:
                chunk = ""

                for char in current:
                    if stringWidth(chunk + char, font, size) <= width:
                        chunk += char
                    else:
                        if chunk:
                            lines.append(chunk)
                        chunk = char

                current = chunk

    if current:
        lines.append(current)

    return lines or [""]


def strip_brackets(value):
    """
    Remove text inside (), {} and [].
    Also remove any remaining bracket symbols.
    """
    value = clean(value)

    value = re.sub(r"\([^)]*\)", "", value)
    value = re.sub(r"\{[^}]*\}", "", value)
    value = re.sub(r"\[[^\]]*\]", "", value)

    return (
        value.replace("(", "")
        .replace(")", "")
        .replace("{", "")
        .replace("}", "")
        .replace("[", "")
        .replace("]", "")
        .strip()
    )


# ============================================================
# MAC ID EXTRACTION
# ============================================================

def extract_mac(value):
    """
    Extract MAC IDs from:
      MAC: 00-21-49-09-74-D7
      MAC ID: 00-21-49-09-74-D7
      00-21-49-09-74-D7
      0021490974D7

    Also supports colon-separated MAC addresses.
    """
    value = clean(value)

    if not value:
        return ""

    # Prefer explicitly labelled MAC addresses.
    labelled = re.search(
        r"\bMAC(?:\s*ID)?\s*[:=]\s*"
        r"((?:[0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}"
        r"|[0-9A-Fa-f]{12})",
        value,
        flags=re.I,
    )

    if labelled:
        return labelled.group(1)

    # Support plain MAC addresses in the dedicated MAC column.
    plain = re.search(
        r"(?<![0-9A-Fa-f])"
        r"((?:[0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}"
        r"|[0-9A-Fa-f]{12})"
        r"(?![0-9A-Fa-f])",
        value,
    )

    if plain:
        return plain.group(1)

    # Preserve existing behaviour for non-MAC values.
    return strip_brackets(value)


def extract_serial_and_mac(serial_value, mac_value):
    """
    Extract the serial number and MAC address independently.

    Input:
        FBE26060499 (MAC: 00-21-49-09-74-D7)

    Output:
        serial = FBE26060499
        mac    = 00-21-49-09-74-D7
    """
    serial = clean(serial_value)
    mac = extract_mac(mac_value)

    # Find a labelled MAC address inside the Serial Number field.
    match = re.search(
        r"\bMAC(?:\s*ID)?\s*[:=]\s*"
        r"((?:[0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}"
        r"|[0-9A-Fa-f]{12})",
        serial,
        flags=re.I,
    )

    if match:
        # Do not overwrite a MAC address already found in the MAC column.
        if not mac:
            mac = match.group(1)

        # Remove only the MAC label and address from the serial text.
        serial = serial[:match.start()] + serial[match.end():]

    # Remove remaining bracketed content from the serial number.
    serial = strip_brackets(serial)

    return serial, mac


# ============================================================
# WINDOWS KEY / OPERATING SYSTEM DETECTION
# ============================================================

WINDOWS_KEY_RE = re.compile(
    r"\b(?:[A-Z0-9]{5}-){4}[A-Z0-9]{5}\b",
    flags=re.I,
)


def extract_windows_key(value):
    """Return a Windows/Product key found anywhere in a cell."""
    text = clean(value)
    if not text:
        return ""
    match = WINDOWS_KEY_RE.search(text)
    return match.group(0) if match else ""


def is_operating_system_text(value):
    text = clean(value).upper()
    if not text:
        return False
    return bool(
        re.search(r"\bWINDOWS?\s*(?:10|11)\b", text)
        or re.search(r"\bWIN\s*(?:10|11)\b", text)
        or "WINDOWS KEY" in text
        or text.startswith("WIN ")
    )


def clean_operating_system(value):
    text = valid_value(value)
    if not text:
        return ""
    return strip_brackets(text)


def intro_labels(item, settings):
    """Build intro labels independently from table visibility/order."""
    candidates = [
        ("model_no", settings.get("model_heading", "MODEL NO"), "intro_model"),
        ("serial_no", settings.get("serial_heading", "SERIAL NO"), "intro_serial"),
        ("mac", settings.get("mac_heading", "MAC ID"), "intro_mac"),
        ("make", settings.get("make_heading", "MAKE"), "intro_make"),
        ("operating_system", settings.get("operating_system_heading", "OPERATING SYSTEM"), "intro_operating_system"),
        ("win_key", settings.get("win_heading", "WINDOW KEY"), "intro_win"),
    ]
    labels = []
    for key, heading, setting_key in candidates:
        if settings.get(setting_key, False) and valid_value(item.get(key)):
            labels.append(heading)
    return labels


def build_intro(item, settings):
    item_name = valid_value(item.get("item")) or "TEMP"
    labels = intro_labels(item, settings)
    intro = f"This test certificate is for {item_name}"
    if labels:
        if len(labels) == 1:
            joined = labels[0]
        elif len(labels) == 2:
            joined = f"{labels[0]} and {labels[1]}"
        else:
            joined = ", ".join(labels[:-1]) + f" and {labels[-1]}"
        intro += " having " + joined
    intro += " as per below table."
    return intro


# ============================================================
# EXCEL READER
# ============================================================

def read_excel(uploaded_file):
    wb = openpyxl.load_workbook(uploaded_file, data_only=True)
    ws = wb[wb.sheetnames[0]]

    company = ""
    date_raw = ""
    header_row = None

    for row in range(1, min(ws.max_row, 20) + 1):
        values = [
            clean(ws.cell(row, col).value).lower().replace(".", "").strip()
            for col in range(1, ws.max_column + 1)
        ]

        for col in range(1, ws.max_column + 1):
            value = clean(ws.cell(row, col).value)
            if value.lower().startswith("company name"):
                company = re.sub(r"^company name\s*:?\s*", "", value, flags=re.I).strip()
            if value.lower().startswith("date"):
                date_raw = re.sub(r"^date\s*:?\s*", "", value, flags=re.I).strip()

        has_name = any(value in values for value in ("name", "item", "material", "item name", "description"))
        has_model_or_serial = any(
            value in values
            for value in (
                "mn", "model", "model no", "model number", "modelno",
                "sn", "serial", "serial no", "serial number", "serialno",
            )
        )
        has_windows = any("windows" in value or value.startswith("win ") for value in values)

        if (has_name and has_model_or_serial) or (has_name and has_windows):
            header_row = row
            break

    if header_row is None:
        raise ValueError(
            "Excel headings could not be found. Add a Name/Item heading along with Model/MN or Serial/SN."
        )

    headers = {}
    for col in range(1, ws.max_column + 1):
        heading = clean(ws.cell(header_row, col).value).lower().replace(".", "").strip()
        if heading:
            headers[heading] = col

    def find_col(*names):
        for name in names:
            if name in headers:
                return headers[name]
        return None

    name_col = find_col("name", "item", "material", "item name", "description")
    make_col = find_col("make", "brand", "manufacturer")
    model_col = find_col("mn", "model", "model no", "model number", "modelno")
    serial_col = find_col("sn", "serial", "serial no", "serial number", "serialno")
    mac_col = find_col("mac", "mac id", "mac address", "macid")
    win_col = find_col("win key", "windows key", "product key", "license key", "licence key")
    os_col = find_col("operating system", "os", "operating system name", "windows version")
    qty_col = find_col("qty", "quantity")

    if not name_col:
        raise ValueError("The Name/Item column is required in Excel.")

    items = []
    previous_item = None

    for row in range(header_row + 1, ws.max_row + 1):
        row_values = [clean(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        row_values_valid = [valid_value(v) for v in row_values]

        name = valid_value(merged_cell_value(ws, row, name_col))
        make = valid_value(merged_cell_value(ws, row, make_col)) if make_col else ""
        model = valid_value(merged_cell_value(ws, row, model_col)) if model_col else ""
        serial = valid_value(ws.cell(row, serial_col).value) if serial_col else ""
        mac = valid_value(ws.cell(row, mac_col).value) if mac_col else ""
        win_key = valid_value(ws.cell(row, win_col).value) if win_col else ""
        operating_system = valid_value(ws.cell(row, os_col).value) if os_col else ""

        # Search the entire row for a Windows key. This handles keys placed
        # in Make, Serial, Model, or any other column.
        detected_key = ""
        for value in row_values_valid:
            detected_key = extract_windows_key(value)
            if detected_key:
                break
        win_key = win_key or detected_key

        serial, mac = extract_serial_and_mac(serial, mac)
        model = strip_brackets(model)

        # Detect an operating-system description anywhere in the row.
        if not operating_system:
            for value in row_values_valid:
                if is_operating_system_text(value) and not extract_windows_key(value):
                    upper = value.upper()
                    if "WINDOWS KEY" not in upper and not WINDOWS_KEY_RE.search(value):
                        operating_system = clean_operating_system(value)
                        break

        # Windows handling has TWO different cases:
        #
        # 1) A real item row named "Windows Key" is its OWN certificate.
        #    Example:
        #    Windows Key | Radhe Info | Win 11 IOT ENT LTSC | KNB2W-...
        #
        # 2) A continuation row whose item/name starts with WIN/WINDOWS and
        #    contains a product key belongs to the PREVIOUS item (usually IPC).
        row_name_upper = name.upper()
        is_standalone_windows_key_item = row_name_upper.strip() in {
            "WINDOWS KEY", "WINDOW KEY", "WINDOWS PRODUCT KEY",
            "WINDOWS LICENSE KEY", "WINDOWS LICENCE KEY",
        }
        same_item_windows_continuation = (
            not is_standalone_windows_key_item
            and previous_item is not None
            and name
            and valid_value(previous_item.get("item")).strip().upper() == row_name_upper.strip()
            and bool(detected_key)
            and (
                is_operating_system_text(model)
                or is_operating_system_text(name)
                or is_operating_system_text(operating_system)
                or is_operating_system_text(serial)
            )
        )

        is_windows_continuation = (
            not is_standalone_windows_key_item
            and (
                row_name_upper.startswith("WIN ")
                or row_name_upper.startswith("WINDOWS ")
                or same_item_windows_continuation
            )
            and bool(detected_key or operating_system or model or serial)
        )

        # A standalone Windows Key row is a real certificate. Move the OS
        # text from MN/model into OPERATING SYSTEM and keep the product key
        # only in WINDOW KEY (not duplicated as SERIAL NO.).
        if is_standalone_windows_key_item:
            if not operating_system and is_operating_system_text(model):
                operating_system = clean_operating_system(model)
            if not operating_system and is_operating_system_text(serial):
                operating_system = clean_operating_system(serial)
            if not win_key:
                win_key = detected_key
            model = ""
            serial = ""

        if is_windows_continuation and previous_item:
            os_text = operating_system
            if not os_text:
                if is_operating_system_text(name) and "WINDOWS KEY" not in row_name_upper:
                    os_text = clean_operating_system(name)
                elif model and is_operating_system_text(model):
                    os_text = clean_operating_system(model)
                elif serial and is_operating_system_text(serial):
                    os_text = clean_operating_system(serial)

            if os_text:
                previous_item["operating_system"] = os_text
            if win_key:
                previous_item["win_key"] = win_key
            previous_item["_windows_detected"] = bool(win_key)
            previous_item["_windows_source_row"] = row
            continue

        # If the key is on the same normal item row, keep it with that item.
        if not any([name, make, model, serial, mac, win_key, operating_system]):
            continue

        if not name:
            if previous_item and (win_key or operating_system or serial or mac):
                if win_key:
                    previous_item["win_key"] = win_key
                if operating_system:
                    previous_item["operating_system"] = operating_system
                if serial and not previous_item.get("serial_no"):
                    previous_item["serial_no"] = serial
                if mac and not previous_item.get("mac"):
                    previous_item["mac"] = mac
                previous_item["_windows_detected"] = bool(win_key)
                previous_item.setdefault("_windows_source_row", row)
                continue
            continue

        qty_raw = ws.cell(row, qty_col).value if qty_col else 1
        try:
            qty = max(1, int(float(qty_raw or 1)))
        except (TypeError, ValueError):
            qty = 1

        item = {
            "item": name,
            "make": make,
            "model_no": model,
            "serial_no": serial,
            "mac": mac,
            "operating_system": operating_system,
            "win_key": win_key,
            "qty": qty,
            "_windows_detected": bool(win_key),
            "_windows_source_row": row if win_key else None,
        }

        items.append(item)
        previous_item = item

    if not items:
        raise ValueError("No item rows were found in the Excel file.")

    company = company or "Customer"
    company = re.sub(r"\s*TVC\s*\d+.*$", "", company, flags=re.I).strip()
    company = re.sub(r"\s*\([^)]*\)\s*$", "", company).strip() or "Customer"

    return {
        "company": company,
        "date": parse_date(date_raw),
        "items": items,
    }


# ============================================================
# DYNAMIC TABLE COLUMNS
# ============================================================

FIELD_DEFINITIONS = {
    "item": ("ITEM", None),
    "make": ("make_heading", "show_make"),
    "model_no": ("model_heading", "show_model"),
    "serial_no": ("serial_heading", "show_serial"),
    "mac": ("mac_heading", "show_mac"),
    "operating_system": ("operating_system_heading", "show_operating_system"),
    "win_key": ("win_heading", "show_win"),
}

DEFAULT_TABLE_ORDER = ["item", "model_no", "serial_no", "operating_system", "win_key", "mac", "make"]


def table_fields(item, settings):
    # ITEM is always first and MAKE is always last.
    # Page Settings controls the order of every other enabled column.
    configured = list(settings.get("table_order") or DEFAULT_TABLE_ORDER)
    middle = [
        key for key in configured
        if key not in ("item", "make")
        and key in FIELD_DEFINITIONS
    ]

    # Keep all known fields available even if an older saved setting does not
    # contain one of them.
    for key in DEFAULT_TABLE_ORDER:
        if key not in ("item", "make") and key not in middle:
            middle.append(key)

    order = ["item"] + middle + ["make"]

    fields = []
    for key in order:
        definition = FIELD_DEFINITIONS.get(key)
        if not definition:
            continue
        heading_key, setting_key = definition
        if key == "item":
            heading = "ITEM"
        else:
            heading = settings.get(heading_key, key.upper())
            if setting_key and not settings.get(setting_key, True):
                continue
            if item.get("_windows_detected"):
                if key in ("model_no", "serial_no") and item.get("_popup_show_model_serial") is False:
                    continue
                if key in ("operating_system", "win_key") and item.get("_popup_show_os_win") is False:
                    continue
        value = valid_value(item.get(key))
        if value:
            fields.append((heading, value))
    return fields


# ============================================================
# PDF GENERATION
# ============================================================

def draw_certificate_page(pdf, item, company, cert_date, settings):
    page_w, page_h = A4

    if settings.get("paper_type") == "Letterhead":
        try:
            logo_record = get_bg_logo()

            if logo_record:
                logo_bytes, _mime_type = logo_record

                pdf.drawImage(
                    ImageReader(io.BytesIO(logo_bytes)),
                    0, 0,
                    width=page_w,
                    height=page_h,
                    preserveAspectRatio=False,
                    mask="auto",
                )
            else:
                pdf.setFont("Helvetica-Bold", 10)
                pdf.drawString(48, page_h - 35, "OptoPixel")

        except Exception as exc:
            raise RuntimeError(
                "Failed to load the Aiven background image. "
                "Check Streamlit Secrets and the database connection."
            ) from exc

    pdf.saveState()

    shift = float(settings.get("top_spacing_mm", 10)) * 72 / 25.4
    pdf.translate(0, -shift)

    left, right = 48, page_w - 48
    content_w = right - left

    fmt = DATE_FORMATS.get(
        settings.get("date_format"), "%d-%b-%y"
    )

    date_text = (
        cert_date.strftime(fmt)
        if isinstance(cert_date, (date, datetime))
        else str(cert_date)
    )

    pdf.setFillColorRGB(0.08, 0.08, 0.08)
    pdf.setFont("Helvetica", 10)
    pdf.drawRightString(right, page_h - 145, "Date: " + date_text)

    title = "TEST CERTIFICATE"
    title_y = page_h - 95

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawCentredString(page_w / 2, title_y, title)

    title_width = stringWidth(title, "Helvetica-Bold", 16)

    pdf.line(
        page_w / 2 - title_width / 2,
        title_y - 3,
        page_w / 2 + title_width / 2,
        title_y - 3,
    )

    concern = "TO WHOM IT MAY CONCERN"
    concern_y = page_h - 190

    pdf.setFont("Helvetica", 12)
    pdf.drawCentredString(page_w / 2, concern_y, concern)

    concern_width = stringWidth(concern, "Helvetica", 12)

    pdf.line(
        page_w / 2 - concern_width / 2,
        concern_y - 3,
        page_w / 2 + concern_width / 2,
        concern_y - 3,
    )

    item_name = valid_value(item.get("item")) or "TEMP"
    fields = table_fields(item, settings)

    intro = build_intro(item, settings)

    intro_lines = wrap_text(
        intro, "Helvetica", 10.5, content_w
    )

    intro_y = page_h - 245
    pdf.setFont("Helvetica", 10.5)

    for index, line in enumerate(intro_lines):
        pdf.drawString(left, intro_y - index * 15, line)

    count = len(fields)

    widths = (
        [content_w]
        if count == 1
        else [content_w * 0.25]
        + [content_w * 0.75 / (count - 1)] * (count - 1)
    )

    wrapped = []
    max_header = 1
    max_value = 1

    for index, (heading, value) in enumerate(fields):
        available = widths[index] - 8

        header_lines = wrap_text(
            heading, "Helvetica-Bold", 8.5, available
        )

        value_lines = wrap_text(
            value, "Helvetica", 8.5, available
        )

        wrapped.append((header_lines, value_lines))

        max_header = max(max_header, len(header_lines))
        max_value = max(max_value, len(value_lines))

    header_h = max(30, max_header * 11 + 12)
    row_h = max(34, max_value * 11 + 12)

    table_top = intro_y - len(intro_lines) * 15 - 25
    table_bottom = table_top - header_h - row_h

    xs = [left]

    for width in widths:
        xs.append(xs[-1] + width)

    pdf.setLineWidth(0.6)

    pdf.line(left, table_top, right, table_top)
    pdf.line(
        left, table_top - header_h, right, table_top - header_h
    )
    pdf.line(left, table_bottom, right, table_bottom)

    for x in xs:
        pdf.line(x, table_top, x, table_bottom)

    for index, (header_lines, value_lines) in enumerate(wrapped):
        center_x = (xs[index] + xs[index + 1]) / 2

        y = (
            table_top
            - (header_h - len(header_lines) * 11) / 2
            - 8
        )

        pdf.setFont("Helvetica-Bold", 8.5)

        for line in header_lines:
            pdf.drawCentredString(center_x, y, line)
            y -= 11

        y = (
            table_top - header_h
            - (row_h - len(value_lines) * 11) / 2
            - 8
        )

        pdf.setFont("Helvetica", 8.5)

        for line in value_lines:
            pdf.drawCentredString(center_x, y, line)
            y -= 11

    y = table_bottom - 35

    statement = (
        "Specified component is tested for all its functionalities "
        "and is performing satisfactorily and found OK."
    )

    pdf.setFont("Helvetica", 10)

    for line in wrap_text(
        statement, "Helvetica", 10, content_w
    ):
        pdf.drawString(left, y, line)
        y -= 15

    y -= 20
    pdf.setFont("Helvetica-Oblique", 10)

    sale_statement = (
        f"The above said item is sold to {company or 'Customer'}."
    )

    for line in wrap_text(
        sale_statement, "Helvetica-Oblique", 10, content_w
    ):
        pdf.drawString(left, y, line)
        y -= 15

    y -= 20
    pdf.setFont("Helvetica", 10)

    pdf.drawString(
        left, y,
        f'For {settings.get("company", DEFAULTS["company"])},'
    )

    pdf.drawString(left, y - 45, "Authorized")

    pdf.restoreState()
    pdf.showPage()


def create_pdf(data, settings):
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)

    for item in data["items"]:
        copies = max(1, int(item.get("qty", 1)))

        item_company = item.get("_company", data["company"])
        item_date = item.get("_date", data["date"])

        for _ in range(copies):
            draw_certificate_page(
                pdf, item, item_company, item_date, settings
            )

    pdf.save()
    return output.getvalue()


# ============================================================
# WORD GENERATION
# ============================================================

def create_word(data, settings):
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    first_page = True

    for item in data["items"]:
        copies = max(1, int(item.get("qty", 1)))

        company = item.get("_company", data["company"])
        cert_date = item.get("_date", data["date"])

        for _ in range(copies):
            if not first_page:
                doc.add_page_break()

            first_page = False

            fmt = DATE_FORMATS.get(
                settings.get("date_format"), "%d-%b-%y"
            )

            date_text = (
                cert_date.strftime(fmt)
                if isinstance(cert_date, (date, datetime))
                else str(cert_date)
            )

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run("Date: " + date_text)

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = p.add_run("TEST CERTIFICATE")
            run.bold = True
            run.font.size = Pt(16)

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = p.add_run("TO WHOM IT MAY CONCERN")
            run.bold = True
            run.font.size = Pt(11)

            item_name = valid_value(item.get("item")) or "TEMP"

            doc.add_paragraph(build_intro(item, settings))

            fields = table_fields(item, settings)

            table = doc.add_table(
                rows=2,
                cols=len(fields),
            )

            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for col_index, (heading, value) in enumerate(fields):
                table.cell(0, col_index).text = heading
                table.cell(1, col_index).text = value

                for row_index in (0, 1):
                    cell = table.cell(row_index, col_index)
                    cell.vertical_alignment = (
                        WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    )

                    for paragraph in cell.paragraphs:
                        paragraph.alignment = (
                            WD_ALIGN_PARAGRAPH.CENTER
                        )

                        for run in paragraph.runs:
                            run.font.size = Pt(8)

                for run in (
                    table.cell(0, col_index)
                    .paragraphs[0].runs
                ):
                    run.bold = True

            doc.add_paragraph(
                "Specified component is tested for all its functionalities "
                "and is performing satisfactorily and found OK."
            )

            doc.add_paragraph(
                f"The above said item is sold to {company}."
            )

            doc.add_paragraph(
                f'For {settings.get("company", DEFAULTS["company"])},'
            )

            doc.add_paragraph("\nAuthorized")

    output = io.BytesIO()
    doc.save(output)

    return output.getvalue()


# ============================================================
# LOCAL WINDOWS AUTO-SAVE
# ============================================================

def save_files_locally(files, company, cert_date):
    if os.name != "nt":
        return None

    date_folder = (
        cert_date.strftime("%d-%b-%y")
        if isinstance(cert_date, (date, datetime))
        else safe_stem(cert_date)
    )

    company_folder = safe_stem(company)

    for root in (r"D:\M-CERTIFICATE", r"C:\M-CERTIFICATE"):
        try:
            os.makedirs(root, exist_ok=True)

            target = os.path.join(
                root, date_folder, company_folder
            )

            os.makedirs(target, exist_ok=True)

            for filename, content in files.items():
                with open(
                    os.path.join(target, filename), "wb"
                ) as handle:
                    handle.write(content)

            return target

        except OSError:
            continue

    return None


# ============================================================
# GENERATE INDIVIDUAL + COMBINED FILES
# ============================================================

def make_entries(parsed, settings):
    entries = []
    counts = {}

    for original in parsed["items"]:
        copies = max(1, int(original.get("qty", 1)))

        for _ in range(copies):
            item = dict(original)

            item["_company"] = item.get(
                "_company", parsed["company"]
            )

            item["_date"] = item.get(
                "_date", parsed["date"]
            )

            stem = item_filename(item)

            counts[stem] = counts.get(stem, 0) + 1

            if counts[stem] > 1:
                stem = f"{stem}_{counts[stem]}"

            one_data = {
                "company": item["_company"],
                "date": item["_date"],
                "items": [item],
            }

            entries.append({
                "label": (
                    f'{item.get("item") or "TEMP"} | '
                    f'Model: {item.get("model_no") or "—"} | '
                    f'Serial: {item.get("serial_no") or "—"}'
                ),
                "stem": stem,
                "pdf_name": stem + ".pdf",
                "docx_name": stem + ".docx",
                "pdf": create_pdf(one_data, settings),
                "docx": create_word(one_data, settings),
                "item": item,
            })

    return entries


def build_bundle(entries, company, cert_date, settings):
    combined_items = []

    for entry in entries:
        item = dict(entry["item"])
        item["qty"] = 1
        combined_items.append(item)

    combined_data = {
        "company": company,
        "date": cert_date,
        "items": combined_items,
    }

    combined_pdf = create_pdf(combined_data, settings)
    combined_word = create_word(combined_data, settings)

    combined_stem = combined_filename(company, cert_date)

    all_files = {}

    for entry in entries:
        all_files[entry["pdf_name"]] = entry["pdf"]
        all_files[entry["docx_name"]] = entry["docx"]

    all_files[combined_stem + ".pdf"] = combined_pdf
    all_files[combined_stem + ".docx"] = combined_word

    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(
        zip_buffer, "w", zipfile.ZIP_DEFLATED
    ) as archive:
        for filename, content in all_files.items():
            archive.writestr(filename, content)

    saved_folder = save_files_locally(
        all_files, company, cert_date
    )

    return {
        "entries": entries,
        "combined_pdf": combined_pdf,
        "combined_docx": combined_word,
        "combined_pdf_name": combined_stem + ".pdf",
        "combined_docx_name": combined_stem + ".docx",
        "zip": zip_buffer.getvalue(),
        "zip_name": combined_stem + "_ALL_FILES.zip",
        "company": company,
        "date": cert_date,
        "saved_folder": saved_folder,
    }


# ============================================================
# DOWNLOADS
# ============================================================

def show_downloads(generated, key_prefix):
    if not generated:
        st.info(
            "First, generate a certificate on the M-SMART PDF "
            "or M-CERTIFICATE page."
        )
        return

    st.subheader("Combined Downloads")

    c1, c2, c3 = st.columns(3)

    with c1:
        st.download_button(
            "⬇ DOWNLOAD ALL — ONE PDF",
            generated["combined_pdf"],
            generated["combined_pdf_name"],
            "application/pdf",
            key=f"{key_prefix}_all_pdf",
            use_container_width=True,
        )

    with c2:
        st.download_button(
            "⬇ DOWNLOAD ALL — ONE WORD",
            generated["combined_docx"],
            generated["combined_docx_name"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key=f"{key_prefix}_all_word",
            use_container_width=True,
        )

    with c3:
        st.download_button(
            "⬇ DOWNLOAD ZIP — ALL FILES",
            generated["zip"],
            generated["zip_name"],
            "application/zip",
            key=f"{key_prefix}_all_zip",
            use_container_width=True,
        )

    st.divider()
    st.subheader("Individual PDF and Word Files")

    for index, entry in enumerate(generated["entries"]):
        with st.container(border=True):
            left, mid, right = st.columns([3, 1, 1])

            with left:
                st.markdown(
                    f"**{entry['item'].get('item') or 'TEMP'}**"
                )
                st.caption(entry["pdf_name"])

            with mid:
                st.download_button(
                    "PDF",
                    entry["pdf"],
                    entry["pdf_name"],
                    "application/pdf",
                    key=f"{key_prefix}_pdf_{index}",
                    use_container_width=True,
                )

            with right:
                st.download_button(
                    "WORD",
                    entry["docx"],
                    entry["docx_name"],
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"{key_prefix}_word_{index}",
                    use_container_width=True,
                )

    st.divider()
    st.subheader("Download Selected Certificates")

    labels = [
        f"{index + 1}. {entry['label']}"
        for index, entry in enumerate(generated["entries"])
    ]

    selected = st.multiselect(
        "Select the certificates you want to download",
        options=labels,
        key=f"{key_prefix}_selected",
    )

    if selected:
        chosen = [
            entry
            for index, entry in enumerate(generated["entries"])
            if labels[index] in selected
        ]

        selected_data = {
            "company": generated["company"],
            "date": generated["date"],
            "items": [],
        }

        for entry in chosen:
            item = dict(entry["item"])
            item["qty"] = 1
            selected_data["items"].append(item)

        selected_pdf = create_pdf(
            selected_data, st.session_state.smart_settings
        )

        selected_word = create_word(
            selected_data, st.session_state.smart_settings
        )

        selected_stem = (
            combined_filename(
                generated["company"], generated["date"]
            )
            + f"_SELECTED_{len(chosen)}"
        )

        c1, c2, c3 = st.columns(3)

        with c1:
            st.download_button(
                "Download Selected — PDF",
                selected_pdf,
                selected_stem + ".pdf",
                "application/pdf",
                key=f"{key_prefix}_selected_pdf",
                use_container_width=True,
            )

        with c2:
            st.download_button(
                "Download Selected — Word",
                selected_word,
                selected_stem + ".docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"{key_prefix}_selected_word",
                use_container_width=True,
            )

        with c3:
            selected_zip = io.BytesIO()

            with zipfile.ZipFile(
                selected_zip, "w", zipfile.ZIP_DEFLATED
            ) as archive:
                for entry in chosen:
                    archive.writestr(
                        entry["pdf_name"], entry["pdf"]
                    )
                    archive.writestr(
                        entry["docx_name"], entry["docx"]
                    )

            st.download_button(
                "Download Selected — ZIP",
                selected_zip.getvalue(),
                selected_stem + ".zip",
                "application/zip",
                key=f"{key_prefix}_selected_zip",
                use_container_width=True,
            )


# ============================================================
# NAVIGATION
# ============================================================

nav_cols = st.columns(4)

pages = [
    "M-SMART PDF",
    "DOWNLOAD PDF WORD",
    "M-CERTIFICATE",
    "PAGE SETTINGS",
]

for col, page_name in zip(nav_cols, pages):
    with col:
        if st.button(
            page_name,
            key="nav_" + page_name,
            type=(
                "primary"
                if st.session_state.page == page_name
                else "secondary"
            ),
            use_container_width=True,
        ):
            st.session_state.page = page_name
            st.rerun()

if st.button("LOG OUT", key="logout_button"):
    st.session_state.authenticated = False
    st.session_state.is_admin = False
    st.rerun()

st.divider()

settings = st.session_state.smart_settings


# ============================================================
# PAGE 1 — EXCEL TO CERTIFICATES
# ============================================================

if st.session_state.page == "M-SMART PDF":
    st.header("M-SMART PDF")

    st.write(
        "Upload Excel files. A separate PDF and editable Word "
        "document will be generated for each item."
    )

    uploads = st.file_uploader(
        "Upload Excel Files (.xlsx)",
        type=["xlsx"],
        accept_multiple_files=True,
        key="excel_upload",
    )

    if uploads:
        try:
            upload_signature = tuple((u.name, getattr(u, "size", 0)) for u in uploads)
            if st.session_state.get("excel_upload_signature") != upload_signature:
                st.session_state.excel_upload_signature = upload_signature
                st.session_state.windows_queue = []
                st.session_state.windows_decisions = {}

            parsed_files = [read_excel(uploaded) for uploaded in uploads]
            all_items = []
            for parsed_file in parsed_files:
                for original in parsed_file["items"]:
                    item = dict(original)
                    item["_company"] = parsed_file["company"]
                    item["_date"] = parsed_file["date"]
                    all_items.append(item)

            company = parsed_files[0]["company"]
            cert_date = parsed_files[0]["date"]
            parsed = {"company": company, "date": cert_date, "items": all_items}

            # Create a warning queue for every imported Windows key.
            if not st.session_state.get("windows_queue"):
                for idx, item in enumerate(all_items):
                    if item.get("_windows_detected") and idx not in st.session_state.get("windows_decisions", {}):
                        st.session_state.windows_queue.append(idx)

            total_certificates = sum(max(1, int(item.get("qty", 1))) for item in all_items)
            st.success(
                f"{len(uploads)} Excel file(s) read successfully. "
                f"Total item rows: {len(all_items)}. Certificates to generate: {total_certificates}."
            )

            # Windows-key popup dialog.
            queue = st.session_state.get("windows_queue", [])
            if queue and hasattr(st, "dialog"):
                idx = queue[0]
                detected_item = all_items[idx]

                @st.dialog("Windows Key Detected")
                def windows_key_dialog():
                    st.warning(f"Sr No {idx + 1} in this data having Windows key?")
                    st.write(f"**Item:** {detected_item.get('item', '')}")
                    if detected_item.get("model_no"):
                        st.write(f"**Model No:** {detected_item.get('model_no')}")
                    if detected_item.get("serial_no"):
                        st.write(f"**Serial No:** {detected_item.get('serial_no')}")
                    if detected_item.get("operating_system"):
                        st.write(f"**Operating System:** {detected_item.get('operating_system')}")
                    st.code(detected_item.get("win_key", ""))

                    st.markdown("### Certificate table options")
                    keep_model_serial = st.checkbox(
                        "Do you want Model No Serial No in table?",
                        value=bool(detected_item.get("model_no") or detected_item.get("serial_no")),
                        key=f"popup_model_serial_{idx}",
                    )
                    keep_os_win = st.checkbox(
                        "Do you want Operating system and Window key in table?",
                        value=bool(detected_item.get("operating_system") or detected_item.get("win_key")),
                        key=f"popup_os_win_{idx}",
                    )

                    c1, c2 = st.columns(2)
                    with c1:
                        if st.button("YES — KEEP KEY", type="primary", use_container_width=True):
                            st.session_state.windows_decisions[idx] = {
                                "keep_key": True,
                                "show_model_serial": keep_model_serial,
                                "show_os_win": keep_os_win,
                            }
                            st.session_state.windows_queue.pop(0)
                            st.rerun()
                    with c2:
                        if st.button("NO — REMOVE KEY", use_container_width=True):
                            st.session_state.windows_decisions[idx] = {
                                "keep_key": False,
                                "show_model_serial": keep_model_serial,
                                "show_os_win": keep_os_win,
                            }
                            st.session_state.windows_queue.pop(0)
                            st.rerun()

                windows_key_dialog()

            # Apply Windows popup decisions before showing the editable grid.
            for decision_idx, decision in st.session_state.get("windows_decisions", {}).items():
                if 0 <= decision_idx < len(all_items):
                    item = all_items[decision_idx]
                    if isinstance(decision, bool):
                        keep_key = decision
                        show_model_serial = True
                        show_os_win = True
                    else:
                        keep_key = bool(decision.get("keep_key", True))
                        show_model_serial = bool(decision.get("show_model_serial", True))
                        show_os_win = bool(decision.get("show_os_win", True))
                    if not keep_key:
                        item["win_key"] = ""
                    item["_popup_show_model_serial"] = show_model_serial
                    item["_popup_show_os_win"] = show_os_win

            st.subheader("Excel Data Preview — Editable")
            editor_rows = []
            for item in all_items:
                editor_rows.append({
                    "Item": item.get("item", ""),
                    "Make": item.get("make", ""),
                    "Model No": item.get("model_no", ""),
                    "Serial No": item.get("serial_no", ""),
                    "Operating System": item.get("operating_system", ""),
                    "WINDOW KEY": item.get("win_key", ""),
                    "MAC ID": item.get("mac", ""),
                    "Qty": int(item.get("qty", 1) or 1),
                })

            edited_df = st.data_editor(
                editor_rows,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                column_config={
                    "Item": st.column_config.TextColumn("ITEM"),
                    "Make": st.column_config.TextColumn("MAKE"),
                    "Model No": st.column_config.TextColumn("MODEL NO"),
                    "Serial No": st.column_config.TextColumn("SERIAL NO"),
                    "Operating System": st.column_config.TextColumn("OPERATING SYSTEM"),
                    "WINDOW KEY": st.column_config.TextColumn("WINDOW KEY"),
                    "MAC ID": st.column_config.TextColumn("MAC ID"),
                    "Qty": st.column_config.NumberColumn("QTY", min_value=1, step=1),
                },
                key="smart_excel_editor",
            )

            if st.button("✨ GENERATE CERTIFICATES", type="primary", use_container_width=True, disabled=bool(st.session_state.get("windows_queue"))):
                updated_items = []

                # st.data_editor() returns a list of dictionaries when the
                # input is a list. Older/newer Streamlit versions may return
                # a pandas DataFrame in other cases, so support both.
                if hasattr(edited_df, "to_dict"):
                    edited_rows = edited_df.to_dict("records")
                else:
                    edited_rows = edited_df

                for idx, row in enumerate(edited_rows):
                    row = dict(row)
                    original = dict(all_items[idx])
                    original["item"] = valid_value(row.get("Item"))
                    original["make"] = valid_value(row.get("Make"))
                    original["model_no"] = strip_brackets(row.get("Model No", ""))
                    original["serial_no"] = valid_value(row.get("Serial No"))
                    original["operating_system"] = valid_value(row.get("Operating System"))
                    original["win_key"] = valid_value(row.get("WINDOW KEY"))
                    original["mac"] = valid_value(row.get("MAC ID"))
                    try:
                        original["qty"] = max(1, int(row.get("Qty", 1) or 1))
                    except (TypeError, ValueError):
                        original["qty"] = 1
                    if all_items[idx].get("_windows_detected"):
                        original["_windows_detected"] = True
                        original["_popup_show_model_serial"] = all_items[idx].get("_popup_show_model_serial", True)
                        original["_popup_show_os_win"] = all_items[idx].get("_popup_show_os_win", True)
                    else:
                        original.pop("_windows_detected", None)
                        original.pop("_windows_source_row", None)
                    updated_items.append(original)

                parsed["items"] = updated_items
                entries = make_entries(parsed, settings)
                st.session_state.smart_generated = build_bundle(entries, company, cert_date, settings)
                st.success(f"{len(entries)} individual certificates generated successfully.")
                saved = st.session_state.smart_generated["saved_folder"]
                if saved:
                    st.info("Files saved locally: " + saved)
                st.rerun()

        except Exception as exc:
            st.error(f"Excel processing error: {exc}")

    generated = st.session_state.smart_generated

    if generated:
        st.divider()
        st.subheader("Generated Certificate Preview")

        st.caption(
            "Review the PDF preview below. "
            "To download files, open the DOWNLOAD PDF WORD page."
        )

        try:
            pdf_doc = fitz.open(
                stream=generated["combined_pdf"],
                filetype="pdf",
            )

            for index, page in enumerate(pdf_doc):
                with st.expander(
                    f"Certificate Preview — Page {index + 1}",
                    expanded=(index == 0),
                ):
                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(1.1, 1.1),
                        alpha=False,
                    )

                    preview_image = Image.open(
                        io.BytesIO(pix.tobytes("png"))
                    )

                    st.image(
                        preview_image,
                        use_container_width=True,
                    )

            pdf_doc.close()

        except Exception as exc:
            st.warning(f"Preview error: {exc}")

        st.dataframe(
            [{
                "Item": entry["item"].get("item", ""),
                "PDF Filename": entry["pdf_name"],
                "Word Filename": entry["docx_name"],
            } for entry in generated["entries"]],
            use_container_width=True,
            hide_index=True,
        )


# ============================================================
# PAGE 2 — DOWNLOADS
# ============================================================

elif st.session_state.page == "DOWNLOAD PDF WORD":
    st.header("DOWNLOAD PDF WORD")

    show_downloads(
        st.session_state.smart_generated,
        "smart",
    )

    generated = st.session_state.smart_generated

    if generated and generated["saved_folder"]:
        st.success(
            "Files saved on your PC: " + generated["saved_folder"]
        )

    elif generated:
        st.info(
            "Local auto-save to D:/C: is available only when Streamlit "
            "runs locally on Windows. Use the download buttons when "
            "running in the cloud."
        )


# ============================================================
# PAGE 3 — MANUAL CERTIFICATE
# ============================================================

elif st.session_state.page == "M-CERTIFICATE":
    st.header("M-CERTIFICATE")

    st.write(
        "Create a certificate manually without Excel. "
        "Empty fields will not appear in the table."
    )

    with st.form("manual_certificate_form"):
        left, right = st.columns(2)

        with left:
            item_name = st.text_input(
                "Item Name",
                placeholder="CAMERA / LENS",
            )

            make = st.text_input("Make")
            model = st.text_input("Model No")
            serial = st.text_input("Serial No")

        with right:
            company = st.text_input(
                "Sold To / Company",
                value="Customer",
            )

            cert_date = st.date_input(
                "Certificate Date",
                value=date.today(),
            )

            mac = st.text_input("MAC ID")
            operating_system = st.text_input("Operating System")
            win_key = st.text_input("WIN KEY")

        st.caption(
            "Model No, Serial No, Make, MAC ID, and WIN KEY are optional."
        )

        submit_manual = st.form_submit_button(
            "GENERATE MANUAL CERTIFICATE",
            type="primary",
            use_container_width=True,
        )

    if submit_manual:
        if not item_name.strip():
            st.error("Item name is required.")

        else:
            company = company.strip() or "Customer"

            # Apply the same serial and MAC extraction logic to manual entries.
            serial_clean, mac_clean = extract_serial_and_mac(
                serial, mac
            )

            item = {
                "item": item_name.strip(),
                "make": valid_value(make),
                "model_no": strip_brackets(model),
                "serial_no": serial_clean,
                "mac": mac_clean,
                "operating_system": valid_value(operating_system),
                "win_key": valid_value(win_key),
                "qty": 1,
                "_company": company,
                "_date": cert_date,
            }

            parsed = {
                "company": company,
                "date": cert_date,
                "items": [item],
            }

            entries = make_entries(parsed, settings)

            st.session_state.manual_generated = build_bundle(
                entries, company, cert_date, settings
            )

            st.success("Manual certificate generated successfully.")
            st.rerun()

    if st.session_state.manual_generated:
        generated = st.session_state.manual_generated

        st.subheader("Manual Certificate Preview")

        try:
            pdf_doc = fitz.open(
                stream=generated["combined_pdf"],
                filetype="pdf",
            )

            page = pdf_doc[0]

            pix = page.get_pixmap(
                matrix=fitz.Matrix(1.2, 1.2),
                alpha=False,
            )

            preview_image = Image.open(
                io.BytesIO(pix.tobytes("png"))
            )

            st.image(
                preview_image,
                use_container_width=False,
            )

            pdf_doc.close()

        except Exception as exc:
            st.warning(f"Preview error: {exc}")

        show_downloads(generated, "manual")


# ============================================================
# PAGE 4 — SETTINGS
# ============================================================

elif st.session_state.page == "PAGE SETTINGS":
    st.header("PAGE SETTINGS")

    st.info(
        "Settings are saved for the current Streamlit session. "
        "Permanent storage after an app restart is not guaranteed."
    )

    settings = st.session_state.smart_settings

    if st.session_state.get("is_admin", False):
        st.subheader("BACKGROUND LOGO — ADMIN ONLY")

        st.caption(
            "The letterhead background image will be stored privately "
            "in the Aiven PostgreSQL database."
        )

        uploaded_bg = st.file_uploader(
            "Upload Letterhead Image (PNG/JPG)",
            type=["png", "jpg", "jpeg"],
            key="admin_bg_logo_upload",
        )

        if uploaded_bg is not None:
            try:
                st.image(
                    uploaded_bg,
                    caption="Uploaded Background Preview",
                    use_container_width=True,
                )
            except Exception:
                pass

            if st.button(
                "SAVE BACKGROUND LOGO TO AIVEN",
                type="primary",
                use_container_width=True,
            ):
                try:
                    save_bg_logo(uploaded_bg)
                    st.success(
                        "Background image saved successfully to the Aiven database."
                    )

                except Exception as exc:
                    st.error(f"Image save failed: {exc}")

        if st.button(
            "CHECK SAVED BACKGROUND LOGO",
            key="check_saved_bg_logo",
        ):
            try:
                saved_logo = get_bg_logo()

                if saved_logo:
                    st.image(
                        saved_logo[0],
                        caption="Current Saved Database Background",
                        use_container_width=True,
                    )
                else:
                    st.warning(
                        "No background image has been uploaded yet."
                    )

            except Exception as exc:
                st.error(f"Database read failed: {exc}")

        st.divider()

    with st.form("settings_form"):
        left, right = st.columns(2)

        with left:
            date_format = st.selectbox(
                "Date Format", list(DATE_FORMATS),
                index=list(DATE_FORMATS).index(settings["date_format"]),
            )
            paper_type = st.selectbox(
                "Paper Type", ["Letterhead", "Blank paper"],
                index=["Letterhead", "Blank paper"].index(settings["paper_type"]),
            )
            top_spacing = st.number_input(
                "Content Top Spacing (mm)", min_value=0, max_value=60,
                value=int(settings["top_spacing_mm"]), step=1,
            )
            issuing_company = st.text_input("Certificate Issued By", value=settings["company"])

        with right:
            make_heading = st.text_input("Make Heading", value=settings["make_heading"])
            model_heading = st.text_input("Model Heading", value=settings["model_heading"])
            serial_heading = st.text_input("Serial Heading", value=settings["serial_heading"])
            mac_heading = st.text_input("MAC Heading", value=settings["mac_heading"])
            os_heading = st.text_input("Operating System Heading", value=settings.get("operating_system_heading", "OPERATING SYSTEM"))
            win_heading = st.text_input("WINDOW KEY Heading", value=settings.get("win_heading", "WINDOW KEY"))

        st.subheader("Optional Table Columns")
        a, b, c, d, e, f = st.columns(6)
        with a:
            show_make = st.checkbox("Show Make", value=settings["show_make"])
        with b:
            show_model = st.checkbox("Show Model No", value=settings["show_model"])
        with c:
            show_serial = st.checkbox("Show Serial No", value=settings["show_serial"])
        with d:
            show_mac = st.checkbox("Show MAC ID", value=settings["show_mac"])
        with e:
            show_os = st.checkbox("Show Operating System", value=settings.get("show_operating_system", True))
        with f:
            show_win = st.checkbox("Show WINDOW KEY", value=settings["show_win"])

        st.subheader("Table Heading Order")
        st.caption("ITEM is always first. MAKE is always last. Set the order of the other columns here.")
        order_options = [
            ("model_no", "MODEL NO"),
            ("serial_no", "SERIAL NO"),
            ("mac", "MAC ID"),
            ("operating_system", "OPERATING SYSTEM"),
            ("win_key", "WINDOW KEY"),
        ]
        option_keys = [x[0] for x in order_options]
        option_labels = dict(order_options)
        current_order = [
            x for x in settings.get("table_order", DEFAULT_TABLE_ORDER)
            if x not in ("item", "make") and x in option_keys
        ]
        for key in option_keys:
            if key not in current_order:
                current_order.append(key)

        position_cols = st.columns(5)
        selected_order = []
        for pos, col in enumerate(position_cols):
            with col:
                current = current_order[pos]
                choices = [k for k in option_keys if k not in selected_order]
                idx = choices.index(current) if current in choices else 0
                chosen = st.selectbox(
                    f"Position {pos + 2}", choices,
                    index=idx,
                    format_func=lambda x: option_labels[x],
                    key=f"table_order_pos_{pos}",
                )
                selected_order.append(chosen)

        st.subheader("Introduction Line — Independent Settings")
        st.caption("These settings control only the sentence above the table, not table columns.")
        i1, i2, i3, i4, i5, i6 = st.columns(6)
        with i1:
            intro_model = st.checkbox("Model No", value=settings.get("intro_model", True))
        with i2:
            intro_serial = st.checkbox("Serial No", value=settings.get("intro_serial", True))
        with i3:
            intro_mac = st.checkbox("MAC ID", value=settings.get("intro_mac", False))
        with i4:
            intro_make = st.checkbox("Make", value=settings.get("intro_make", False))
        with i5:
            intro_os = st.checkbox("Operating System", value=settings.get("intro_operating_system", False))
        with i6:
            intro_win = st.checkbox("WINDOW KEY", value=settings.get("intro_win", False))

        save_settings = st.form_submit_button("SAVE SETTINGS", type="primary", use_container_width=True)

    if save_settings:
        st.session_state.smart_settings.update({
            "date_format": date_format,
            "paper_type": paper_type,
            "top_spacing_mm": int(top_spacing),
            "company": issuing_company.strip() or DEFAULTS["company"],
            "make_heading": make_heading.strip() or "MAKE",
            "model_heading": model_heading.strip() or "MODEL NO",
            "serial_heading": serial_heading.strip() or "SERIAL NO",
            "mac_heading": mac_heading.strip() or "MAC ID",
            "operating_system_heading": os_heading.strip() or "OPERATING SYSTEM",
            "win_heading": win_heading.strip() or "WINDOW KEY",
            "show_make": show_make,
            "show_model": show_model,
            "show_serial": show_serial,
            "show_mac": show_mac,
            "show_operating_system": show_os,
            "show_win": show_win,
            "table_order": ["item"] + selected_order + ["make"],
            "intro_model": intro_model,
            "intro_serial": intro_serial,
            "intro_mac": intro_mac,
            "intro_make": intro_make,
            "intro_operating_system": intro_os,
            "intro_win": intro_win,
        })
        st.session_state.smart_generated = None
        st.session_state.manual_generated = None
        st.success("Settings saved for this session.")
        st.rerun()

    if st.button("RESTORE DEFAULT SETTINGS"):
        st.session_state.smart_settings = DEFAULTS.copy()
        st.session_state.smart_generated = None
        st.session_state.manual_generated = None
        st.rerun()


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.markdown(
    '<div style="text-align:center;color:#728096;font-size:12px">'
    'M-CERTIFICATE · SMART CERTIFICATE STUDIO '
    'MADE BY ~ MEET MEVADA'
    '</div>',
    unsafe_allow_html=True,
        )
