import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import streamlit as st
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.lib.utils import ImageReader


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="M-CERTIFICATE",
    page_icon="🖨️",
    layout="wide",
)

st.title("M-CERTIFICATE")
st.caption("PRINT | SCANNER • Certificate PDF + Word • Linux CUPS/SANE")


# ============================================================
# SAVE LOCATION
# ============================================================

def get_mcertificate_root():
    """
    Linux server storage.
    Files saved here belong to the SERVER, not the user's PC.
    """
    preferred = Path("/home/appuser/M-CERTIFICATE")

    try:
        preferred.mkdir(parents=True, exist_ok=True)
        test_file = preferred / ".write_test"
        test_file.write_text("ok", encoding="utf-8")
        test_file.unlink(missing_ok=True)
        return preferred
    except OSError:
        fallback = Path.home() / "M-CERTIFICATE"
        fallback.mkdir(parents=True, exist_ok=True)
        return fallback


def safe_folder_name(name, default="Customer"):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "", str(name or "").strip())
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name or default


def safe_filename(name, default="document.pdf"):
    name = Path(str(name or "")).name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "", name).strip(" .")
    if not name:
        name = default
    if not name.lower().endswith(".pdf"):
        name += ".pdf"
    return name


def get_certificate_folder(cert_date, sold_to):
    root = get_mcertificate_root()
    date_folder = cert_date.strftime("%d-%b-%y")
    company_folder = safe_folder_name(sold_to)
    folder = root / date_folder / company_folder
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def save_bytes(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


# ============================================================
# WORD DOCUMENT
# ============================================================

def certificate_word_bytes(data):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise RuntimeError(
            "Missing python-docx. Add python-docx to requirements.txt."
        ) from exc

    doc = Document()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Date: " + data["date"].strftime("%d-%b-%y"))
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TEST CERTIFICATE")
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TO WHOM IT MAY CONCERN")
    r.font.size = Pt(12)

    labels = [label for _, label in data["selected_columns"]]

    if not labels:
        intro = f"This test certificate is for {data['item_name']} as per below table"
    elif len(labels) == 1:
        intro = (
            f"This test certificate is for {data['item_name']} "
            f"having {labels[0]} as per below table"
        )
    elif len(labels) == 2:
        intro = (
            f"This test certificate is for {data['item_name']} "
            f"having {labels[0]} and {labels[1]} as per below table"
        )
    else:
        intro = (
            f"This test certificate is for {data['item_name']} "
            f"having {', '.join(labels[:-1])} and {labels[-1]} "
            "as per below table"
        )

    doc.add_paragraph(intro)

    columns = [("item_name", "ITEM", data["item_name"])]
    for key, label in data["selected_columns"]:
        columns.append((key, label, data[key]))

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"

    for i, (_, label, _) in enumerate(columns):
        table.rows[0].cells[i].text = str(label)

    row = table.add_row()
    for i, (_, _, value) in enumerate(columns):
        row.cells[i].text = str(value)

    doc.add_paragraph()
    doc.add_paragraph(
        "Specified component is tested for all its functionalities "
        "and is performing satisfactorily and found OK."
    )
    doc.add_paragraph(f'The above said item is sold to {data["sold_to"]}.')
    doc.add_paragraph()
    doc.add_paragraph("For OptoPixel Pvt. Ltd.,")
    doc.add_paragraph()
    doc.add_paragraph("Authorized")

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


# ============================================================
# CERTIFICATE PDF
# ============================================================

def wrap_words(text, font, size, max_width):
    words = str(text).split()
    lines = []
    line = ""

    for word in words:
        test = word if not line else line + " " + word

        if stringWidth(test, font, size) <= max_width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word

    if line:
        lines.append(line)

    return lines


def make_certificate_pdf(data):
    output = io.BytesIO()
    c = canvas.Canvas(output, pagesize=A4)
    width, height = A4

    c.setFont("Helvetica", 12)
    c.drawRightString(
        width - 50,
        height - 100,
        "Date: " + data["date"].strftime("%d-%b-%y"),
    )

    title = "TEST CERTIFICATE"
    c.setFont("Helvetica-Bold", 16)
    title_y = height - 145
    c.drawCentredString(width / 2, title_y, title)

    title_width = stringWidth(title, "Helvetica-Bold", 16)
    c.line(
        width / 2 - title_width / 2,
        title_y - 3,
        width / 2 + title_width / 2,
        title_y - 3,
    )

    subtitle = "TO WHOM IT MAY CONCERN"
    c.setFont("Helvetica", 12)
    subtitle_y = height - 220
    c.drawCentredString(width / 2, subtitle_y, subtitle)

    subtitle_width = stringWidth(subtitle, "Helvetica", 12)
    c.line(
        width / 2 - subtitle_width / 2,
        subtitle_y - 3,
        width / 2 + subtitle_width / 2,
        subtitle_y - 3,
    )

    labels = [label for _, label in data["selected_columns"]]
    item_name = data["item_name"]

    if not labels:
        intro = f"This test certificate is for {item_name} as per below table"
    elif len(labels) == 1:
        intro = f"This test certificate is for {item_name} having {labels[0]} as per below table"
    elif len(labels) == 2:
        intro = f"This test certificate is for {item_name} having {labels[0]} and {labels[1]} as per below table"
    else:
        intro = (
            f"This test certificate is for {item_name} having "
            f"{', '.join(labels[:-1])} and {labels[-1]} as per below table"
        )

    c.setFont("Helvetica", 11)
    intro_y = height - 260

    for i, line in enumerate(wrap_words(intro, "Helvetica", 11, 490)[:3]):
        c.drawString(50, intro_y - i * 16, line)

    columns = [("item_name", "ITEM", data["item_name"])]
    for key, label in data["selected_columns"]:
        columns.append((key, label, data[key]))

    left = 50
    right = width - 50
    table_top = height - 330
    header_h = 32
    row_h = 35
    table_bottom = table_top - header_h - row_h
    col_width = (right - left) / len(columns)

    xs = [left + i * col_width for i in range(len(columns) + 1)]

    c.setLineWidth(0.6)
    for x in xs:
        c.line(x, table_top, x, table_bottom)

    for y in (table_top, table_top - header_h, table_bottom):
        c.line(left, y, right, y)

    for i, (_, label, _) in enumerate(columns):
        x_center = (xs[i] + xs[i + 1]) / 2
        c.setFont("Helvetica-Bold", 9)
        c.drawCentredString(x_center, table_top - 20, str(label)[:35])

    for i, (_, _, value) in enumerate(columns):
        x_center = (xs[i] + xs[i + 1]) / 2
        max_width = xs[i + 1] - xs[i] - 8
        text = str(value)
        font_size = 9

        while font_size > 6 and stringWidth(text, "Helvetica", font_size) > max_width:
            font_size -= 0.25

        c.setFont("Helvetica", font_size)
        c.drawCentredString(x_center, table_bottom + 12, text)

    body_y = height - 415
    c.setFont("Helvetica", 10.5)
    c.drawString(
        50,
        body_y,
        "Specified component is tested for all its functionalities and is performing",
    )
    c.drawString(50, body_y - 17, "satisfactorily and found OK.")

    sold_text = f'The above said item is sold to {data["sold_to"]}.'
    c.setFont("Helvetica-Oblique", 10.5)

    for i, line in enumerate(wrap_words(sold_text, "Helvetica-Oblique", 10.5, 490)[:3]):
        c.drawString(50, height - 480 - i * 16, line)

    c.setFont("Helvetica", 10.5)
    c.drawString(50, height - 550, "For OptoPixel Pvt. Ltd.,")
    c.drawString(50, height - 600, "Authorized")

    c.showPage()
    c.save()
    return output.getvalue()


# ============================================================
# SAVE CERTIFICATE + CREATE DOWNLOAD ZIP
# ============================================================

def save_certificate_files(data, pdf_bytes, filename_base):
    folder = get_certificate_folder(data["date"], data["sold_to"])
    base = safe_folder_name(filename_base, "test_certificate")
    base = os.path.splitext(base)[0]

    pdf_path = save_bytes(folder / f"{base}.pdf", pdf_bytes)
    word_bytes = certificate_word_bytes(data)
    docx_path = save_bytes(folder / f"{base}.docx", word_bytes)

    return folder, pdf_path, docx_path, word_bytes


def create_download_zip(files, date_folder, company_folder):
    """
    files = [(filename, bytes), ...]
    ZIP contains:
    M-CERTIFICATE/date/company/filename
    """
    output = io.BytesIO()
    prefix = f"M-CERTIFICATE/{date_folder}/{company_folder}"

    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
        for filename, file_bytes in files:
            zf.writestr(f"{prefix}/{filename}", file_bytes)

    output.seek(0)
    return output.getvalue()


# ============================================================
# LINUX CUPS PRINTER HELPERS
# ============================================================

def get_linux_printers():
    """
    Returns printers visible to the Linux server's CUPS system.
    This does NOT list printers attached to the user's Windows PC.
    """
    if not shutil.which("lpstat"):
        return [], "CUPS command 'lpstat' is not installed on this server."

    try:
        result = subprocess.run(
            ["lpstat", "-e"],
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )
        if result.returncode != 0:
            return [], result.stderr.strip() or "Could not query CUPS."

        return [line.strip() for line in result.stdout.splitlines() if line.strip()], None

    except Exception as exc:
        return [], str(exc)


def print_pdf_linux(pdf_bytes, printer_name, copies=1):
    if not shutil.which("lp"):
        raise RuntimeError("CUPS printing command 'lp' is not installed.")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as temp:
        temp.write(pdf_bytes)
        temp_path = temp.name

    try:
        result = subprocess.run(
            [
                "lp",
                "-d", printer_name,
                "-n", str(int(copies)),
                "-o", "media=A4",
                temp_path,
            ],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )

        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or "CUPS printing failed.")

        return result.stdout.strip() or "Print job submitted to CUPS."

    finally:
        try:
            os.unlink(temp_path)
        except OSError:
            pass


# ============================================================
# LINUX SANE SCANNER HELPERS
# ============================================================

def get_linux_scanners():
    """
    List scanners visible to SANE on the Linux server.
    """
    if not shutil.which("scanimage"):
        return [], "SANE command 'scanimage' is not installed on this server."

    result = subprocess.run(
        ["scanimage", "-L"],
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
    )

    combined = (result.stdout or "") + "\n" + (result.stderr or "")
    devices = []

    for line in combined.splitlines():
        match = re.search(r"device [`'](.+?)[`'] is a (.+)", line.strip())
        if match:
            devices.append({
                "device": match.group(1),
                "name": match.group(2),
            })

    if not devices:
        return [], combined.strip() or "No SANE scanners found."

    return devices, None


def scan_one_page_sane(device_name):
    if not shutil.which("scanimage"):
        raise RuntimeError("SANE 'scanimage' command is not installed.")

    with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as temp:
        temp_path = temp.name

    try:
        result = subprocess.run(
            [
                "scanimage",
                "-d", device_name,
                "--format=jpeg",
                "--resolution", "150",
                "-o", temp_path,
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )

        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or "SANE scan failed.")

        image_bytes = Path(temp_path).read_bytes()

        if not image_bytes:
            raise RuntimeError("Scanner returned an empty image.")

        with Image.open(io.BytesIO(image_bytes)) as im:
            im.verify()

        return image_bytes

    finally:
        try:
            os.unlink(temp_path)
        except OSError:
            pass


def make_scanned_pdf(pages):
    if not pages:
        return None

    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    page_width, page_height = A4
    margin = 20

    for page_bytes in pages:
        with Image.open(io.BytesIO(page_bytes)) as original:
            image = original.convert("RGB")
            img_width, img_height = image.size

            scale = min(
                (page_width - margin * 2) / img_width,
                (page_height - margin * 2) / img_height,
            )

            draw_width = img_width * scale
            draw_height = img_height * scale
            x = (page_width - draw_width) / 2
            y = (page_height - draw_height) / 2

            pdf.drawImage(
                ImageReader(image),
                x, y,
                width=draw_width,
                height=draw_height,
                preserveAspectRatio=True,
            )
            pdf.showPage()

    pdf.save()
    output.seek(0)
    return output.getvalue()


# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "mcert_pdf": None,
    "mcert_word": None,
    "mcert_data": None,
    "mcert_copies": 1,
    "scan_pages": [],
    "scan_pdf": None,
    "printer_connected": True,
    "available_scanners": [],
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# MAIN MENU
# ============================================================

menu = st.radio(
    "MENU",
    ["PRINT", "SCANNER"],
    horizontal=True,
    key="main_menu",
)

st.divider()


# ============================================================
# PRINT PAGE
# ============================================================

if menu == "PRINT":
    st.header("🖨️ PRINT")
    st.caption("Generate a certificate, save PDF + Word, or print via Linux CUPS.")

    left, right = st.columns(2)

    with left:
        date_text = st.text_input(
            "Date (DD-MMM-YY)",
            value=datetime.now().strftime("%d-%b-%y"),
        )
        item_name = st.text_input("Item Name")
        show_model = st.checkbox("Model No.")
        show_serial = st.checkbox("Serial No.")
        show_make = st.checkbox("Make")

    with right:
        model_no = st.text_input("Model No. Value") if show_model else ""
        serial_no = st.text_input("Serial No. Value") if show_serial else ""
        make = st.text_input("Make Value") if show_make else ""
        sold_to = st.text_input("Sold To / Company Name")

    selected_columns = []
    if show_model:
        selected_columns.append(("model_no", "Model No."))
    if show_serial:
        selected_columns.append(("serial_no", "Serial No."))
    if show_make:
        selected_columns.append(("make", "Make"))

    copies = st.number_input(
        "Number of Print Copies",
        min_value=1,
        max_value=20,
        value=1,
        step=1,
    )

    if st.button("📄 GENERATE CERTIFICATE", type="primary", use_container_width=True):
        try:
            parsed_date = datetime.strptime(date_text.strip(), "%d-%b-%y").date()
        except ValueError:
            parsed_date = None

        required = {
            "Item Name": item_name,
            "Sold To": sold_to,
        }
        if show_model:
            required["Model No."] = model_no
        if show_serial:
            required["Serial No."] = serial_no
        if show_make:
            required["Make"] = make

        missing = [name for name, value in required.items() if not value.strip()]

        if parsed_date is None:
            st.error("Date must be DD-MMM-YY, for example 26-Oct-26.")
        elif missing:
            st.error("Please fill in: " + ", ".join(missing))
        else:
            data = {
                "date": parsed_date,
                "item_name": item_name.strip(),
                "model_no": model_no.strip(),
                "serial_no": serial_no.strip(),
                "make": make.strip(),
                "sold_to": sold_to.strip(),
                "selected_columns": selected_columns,
            }

            try:
                pdf_bytes = make_certificate_pdf(data)
                word_bytes = certificate_word_bytes(data)

                st.session_state["mcert_pdf"] = pdf_bytes
                st.session_state["mcert_word"] = word_bytes
                st.session_state["mcert_data"] = data
                st.session_state["mcert_copies"] = int(copies)

                st.success("Certificate generated successfully.")

            except Exception as exc:
                st.error(f"Certificate generation failed: {exc}")

    if st.session_state.get("mcert_pdf") and st.session_state.get("mcert_data"):
        st.divider()
        st.subheader("💾 SAVE / DOWNLOAD CERTIFICATE")

        data = st.session_state["mcert_data"]
        filename_base = st.text_input("File Name", value="test_certificate")
        filename_base = safe_folder_name(filename_base, "test_certificate")
        filename_base = os.path.splitext(filename_base)[0]

        pdf_bytes = st.session_state["mcert_pdf"]
        word_bytes = st.session_state["mcert_word"]

        if st.button("💾 SAVE PDF + WORD ON SERVER", use_container_width=True):
            try:
                folder, pdf_path, docx_path, _ = save_certificate_files(
                    data, pdf_bytes, filename_base
                )
                st.success("Saved on Linux server.")
                st.code(str(folder))
            except Exception as exc:
                st.error(f"Save failed: {exc}")

        st.download_button(
            "⬇️ DOWNLOAD PDF",
            data=pdf_bytes,
            file_name=filename_base + ".pdf",
            mime="application/pdf",
            use_container_width=True,
        )

        st.download_button(
            "⬇️ DOWNLOAD WORD",
            data=word_bytes,
            file_name=filename_base + ".docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

        company_folder = safe_folder_name(data["sold_to"])
        date_folder = data["date"].strftime("%d-%b-%y")

        zip_bytes = create_download_zip(
            [
                (filename_base + ".pdf", pdf_bytes),
                (filename_base + ".docx", word_bytes),
            ],
            date_folder,
            company_folder,
        )

        st.download_button(
            "📦 DOWNLOAD COMPANY + DATE FOLDER (ZIP)",
            data=zip_bytes,
            file_name=f"M-CERTIFICATE_{date_folder}_{company_folder}.zip",
            mime="application/zip",
            type="primary",
            use_container_width=True,
        )

        st.divider()
        st.subheader("🖨️ Linux CUPS Printer")

        if st.session_state["printer_connected"]:
            printers, error = get_linux_printers()

            if error:
                st.warning(error)
            elif not printers:
                st.warning("No printer is visible to this Linux server.")
            else:
                selected_printer = st.selectbox("Select server printer", printers)

                if st.button("🖨️ PRINT CERTIFICATE NOW", use_container_width=True):
                    try:
                        message = print_pdf_linux(
                            pdf_bytes,
                            selected_printer,
                            st.session_state["mcert_copies"],
                        )
                        st.success(message)
                    except Exception as exc:
                        st.error(f"Printing failed: {exc}")
        else:
            st.warning("Printer workflow is disconnected in this app.")

else:
    # ========================================================
    # SCANNER PAGE
    # ========================================================

    st.header("📷 SCANNER")
    st.caption("Linux SANE scanner • Multi-page PDF")

    st.subheader("Printer Connection")

    if st.session_state["printer_connected"]:
        st.success("Printer workflow: ON")

        if st.button("🔌 DISCONNECT PRINTER", use_container_width=True):
            st.session_state["printer_connected"] = False
            st.rerun()
    else:
        st.warning("Printer workflow: OFF")

        if st.button("🔌 RECONNECT PRINTER", use_container_width=True):
            st.session_state["printer_connected"] = True
            st.rerun()

    st.divider()

    if st.button("🔄 FIND SCANNERS", use_container_width=True):
        scanners, error = get_linux_scanners()
        st.session_state["available_scanners"] = scanners
        st.session_state["scanner_error"] = error

    scanners = st.session_state["available_scanners"]

    if st.session_state.get("scanner_error"):
        st.warning(st.session_state["scanner_error"])

    if scanners:
        selected = st.selectbox(
            "Select Scanner",
            scanners,
            format_func=lambda x: x["name"],
        )

        count = len(st.session_state["scan_pages"])
        st.info(f"{count} page(s) scanned.")

        button_label = "📷 START SCAN" if count == 0 else "➕ SCAN NEXT PAGE"

        if st.button(button_label, type="primary", use_container_width=True):
            with st.spinner("Scanning page from the Linux server..."):
                try:
                    page = scan_one_page_sane(selected["device"])
                    st.session_state["scan_pages"].append(page)
                    st.session_state["scan_pdf"] = make_scanned_pdf(
                        st.session_state["scan_pages"]
                    )
                    st.rerun()
                except Exception as exc:
                    st.error(f"Scan failed: {exc}")

    else:
        st.info("Click FIND SCANNERS to check scanners connected to the Linux server.")

    st.divider()
    st.subheader("📤 Upload scanned images (optional)")

    uploaded_images = st.file_uploader(
        "Upload scanned JPG / PNG pages to combine into one PDF",
        type=["jpg", "jpeg", "png"],
        accept_multiple_files=True,
    )

    if uploaded_images and st.button("ADD UPLOADED PAGES", use_container_width=True):
        for uploaded in uploaded_images:
            try:
                image_bytes = uploaded.getvalue()
                with Image.open(io.BytesIO(image_bytes)) as im:
                    im.verify()
                st.session_state["scan_pages"].append(image_bytes)
            except Exception:
                st.warning(f"Could not read image: {uploaded.name}")

        st.session_state["scan_pdf"] = make_scanned_pdf(
            st.session_state["scan_pages"]
        )
        st.rerun()

    pages = st.session_state["scan_pages"]

    if pages:
        st.success(f"{len(pages)} page(s) ready.")

        for i, page in enumerate(pages):
            try:
                st.image(page, caption=f"Page {i + 1}", use_container_width=True)
            except Exception:
                st.caption(f"Page {i + 1}")

        if st.button("🗑️ CLEAR ALL SCANS", use_container_width=True):
            st.session_state["scan_pages"] = []
            st.session_state["scan_pdf"] = None
            st.rerun()

        scan_company = st.text_input("Sold To / Company Name", key="scan_company")
        scan_filename = safe_filename(
            st.text_input("PDF File Name", value="scanned_document.pdf"),
            "scanned_document.pdf",
        )

        scan_pdf = st.session_state["scan_pdf"]

        if scan_pdf:
            st.download_button(
                "⬇️ DOWNLOAD SCANNED PDF",
                data=scan_pdf,
                file_name=scan_filename,
                mime="application/pdf",
                type="primary",
                use_container_width=True,
            )

            if scan_company.strip():
                scan_date_folder = datetime.now().strftime("%d-%b-%y")
                company_folder = safe_folder_name(scan_company)

                scan_zip = create_download_zip(
                    [(scan_filename, scan_pdf)],
                    scan_date_folder,
                    company_folder,
                )

                st.download_button(
                    "📦 DOWNLOAD DATE + COMPANY FOLDER (ZIP)",
                    data=scan_zip,
                    file_name=f"M-CERTIFICATE_{scan_date_folder}_{company_folder}.zip",
                    mime="application/zip",
                    use_container_width=True,
                )

                if st.button("💾 SAVE SCANNED PDF ON SERVER", use_container_width=True):
                    try:
                        folder = get_certificate_folder(
                            datetime.now().date(),
                            scan_company,
                        )
                        saved_path = save_bytes(folder / scan_filename, scan_pdf)
                        st.success("Scanned PDF saved on Linux server.")
                        st.code(str(saved_path))
                    except Exception as exc:
                        st.error(f"Save failed: {exc}")

st.divider()
st.caption("M-CERTIFICATE • PRINT | SCANNER • Linux CUPS/SANE")
